# -*- coding: utf-8 -*-
"""図版のデザインシステム（PowerPoint）。

`tools/figures/build.py` から使う。1スライド = 1図版、キャンバスは 1200x630px 固定。

⚠️ **書き出しは LibreOffice ヘッドレス。** PowerPoint で開いたときの見た目ではなく、
LibreOffice が描いた結果が本番に載る。効果を足したら必ず PNG を見て確かめること。
"""
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

W, H = 1200, 630

PX = lambda v: Emu(int(round(v * 914400 / 96)))   # 96dpi
PT = lambda px: Pt(px * 0.75)                     # 1px @96dpi = 0.75pt

# 色はサイトの styles.css と揃える
BG    = RGBColor(0xFC, 0xFC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK   = RGBColor(0x2E, 0x2E, 0x33)
MUTE  = RGBColor(0x6A, 0x71, 0x7B)
NAVY  = RGBColor(0x01, 0x41, 0x72)
NAVY2 = RGBColor(0x0B, 0x5C, 0x99)
PALE  = RGBColor(0xDA, 0xE7, 0xF2)
PALE2 = RGBColor(0xEE, 0xF4, 0xF9)
BAND  = RGBColor(0xF2, 0xF5, 0xF8)
HILI  = RGBColor(0xE6, 0xEE, 0xF6)
RULE  = RGBColor(0xE2, 0xE7, 0xEC)
WARN  = RGBColor(0xA8, 0x41, 0x0F)
WARNB = RGBColor(0xFD, 0xF2, 0xE4)

FONT = "Yu Gothic UI"

# 版面（左右の余白と本文の幅）
M = 64
INNER = W - M * 2


def est_width(body, size):
    """文字列の描画幅のおおよそ。**折り返しを事前に捕まえるため。**

    ⚠️ `word_wrap = False` を LibreOffice が尊重しない。枠より長い見出しは黙って
    2行になり、下の小見出しに重なる（2026-09-01 に kyatatsu-secchi で実際に起きた）。
    全角＝size、半角＝size*0.55 の粗い近似で足りる（判定は「収まるか」だけなので）。
    """
    w = 0.0
    for ch in body:
        w += size if ord(ch) > 0x2E80 else size * 0.55
    return w


class Fig:
    """1枚の図版。すべての座標は 1200x630 の px で書く。"""

    def __init__(self, slide):
        self.s = slide
        # 実際に描いた文字列を全部ためる。build.py がここから数値を拾って
        # 元の SVG と突き合わせる（**渡した値ではなく、画像に出た値**を照合する）
        self.drawn = []
        # 文字の実描画範囲。build.py が総当たりで重なりとはみ出しを検査する
        self.boxes = []
        self.rect(0, 0, W, H, BG)
        self.rect(0, 0, 12, H, NAVY)          # 左のアクセント帯

    # ---------------------------------------------------------------- 部品

    def _plain(self, sh):
        """テーマ既定の効果を外す。

        ⚠️ `shadow.inherit = False` だけでは足りない。`add_shape` が付ける `<p:style>` の
        `effectRef` をテーマから解決して、**LibreOffice が既定の影を描いてしまう**
        （2026-09-01 に実測。行の帯と罫線に影が出た）。`<p:style>` ごと外す。
        """
        sh.shadow.inherit = False
        st = sh._element.find(qn("p:style"))
        if st is not None:
            sh._element.remove(st)
        return sh

    def shadow(self, sh, blur=16, dist=3, alpha=0.14):
        """python-pptx に影を付ける API が無いので spPr に直接入れる"""
        spPr = sh._element.spPr
        old = spPr.find(qn("a:effectLst"))
        if old is not None:
            spPr.remove(old)
        spPr.append(parse_xml(
            '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:outerShdw blurRad="{int(blur * 12700)}" dist="{int(dist * 12700)}"'
            ' dir="5400000" rotWithShape="0">'
            f'<a:srgbClr val="0B2B45"><a:alpha val="{int(alpha * 100000)}"/></a:srgbClr>'
            "</a:outerShdw></a:effectLst>"))
        return sh

    def ghost(self, x, y, w, h, line=None):
        """破線の輪郭だけ。「広げるとここまで」を面で示すのに使う"""
        sh = self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, PX(x), PX(y), PX(w), PX(h))
        sh.fill.background()
        sh.line.color.rgb = line or NAVY
        sh.line.width = PX(2)
        sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return self._plain(sh)

    def rect(self, x, y, w, h, fill, line=None, lw=2, radius=None, shape=None):
        shp = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE)
        sh = self.s.shapes.add_shape(shp, PX(x), PX(y), PX(w), PX(h))
        if radius is not None and shp == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = PX(lw)
        return self._plain(sh)

    def oval(self, x, y, w, h, fill, line=None, lw=2):
        sh = self.s.shapes.add_shape(MSO_SHAPE.OVAL, PX(x), PX(y), PX(w), PX(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = PX(lw)
        return self._plain(sh)

    def bar(self, x, y, w, h, c1, c2, line=None, lw=2, radius=0.22):
        """棒。上から下へのグラデーションで、平面よりも面が立つ"""
        sh = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PX(x), PX(y), PX(max(w, 1)), PX(h))
        sh.adjustments[0] = radius
        sh.fill.gradient()
        sh.fill.gradient_angle = 90.0
        st = sh.fill.gradient_stops
        st[0].color.rgb, st[0].position = c1, 0.0
        st[1].color.rgb, st[1].position = c2, 1.0
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = PX(lw)
        return self._plain(sh)

    def rule(self, x, y, w, color=RULE):
        return self.rect(x, y, w, 1, color)

    def text(self, x, y, w, h, body, size, color, bold=False, align=PP_ALIGN.LEFT, wrap=False):
        return self.rich(x, y, w, h, [(body, size, color, bold)], align, wrap)

    def rich(self, x, y, w, h, parts, align=PP_ALIGN.LEFT, wrap=False):
        """1行の中でサイズや色を変える（「14.5」＋「倍」を1つの塊として置く）"""
        tb = self.s.shapes.add_textbox(PX(x), PX(y), PX(w), PX(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = align
        cx = x
        for body, size, color, bold in parts:
            self.drawn.append(body)
            bw = est_width(body, size)
            bh = size * 1.15
            if align == PP_ALIGN.RIGHT:
                self.boxes.append((x + w - est_width("".join(p[0] for p in parts), size),
                                   y + h / 2 - bh / 2, bw, bh, body))
            else:
                self.boxes.append((cx, y + h / 2 - bh / 2, bw, bh, body))
            cx += bw
            r = p.add_run()
            r.text = body
            r.font.size = PT(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = FONT
        return tb

    # ---------------------------------------------------------------- 定型

    def header(self, title, sub, key=None, key_label=None, key_unit="", band=True):
        """見出し帯。**結論を図の主役にする。**

        紺のベタ帯に白抜きの見出しを置き、右に「この図で言いたい数字ひとつ」を
        白いカードで抜く。地の色と反転するので、縮小表示でも最初に目に入る。
        （`specs/article-rubric.md`「eyecatch を単体で見るだけで記事の結論が分かる」＝5点の条件）
        """
        tw = 780 if key else INNER
        for body, size, label in ((title, 40, "見出し"), (sub, 19, "小見出し")):
            got = est_width(body, size)
            if got > tw:
                raise ValueError(
                    f"{label}が枠に収まりません（推定 {got:.0f}px > {tw}px）: {body}"
                    " / 対処: 短くするか、key を外して横幅を広げる。"
                    " そのままだと LibreOffice が黙って折り返して下の行に重なります")
        if band:
            self.rect(0, 0, W, 158, NAVY)
            self.text(M, 38, tw, 50, title, 40, WHITE, bold=True)
            self.text(M, 92, tw, 26, sub, 19, PALE)
            if key:
                card = self.rect(872, 26, 264, 106, WHITE, radius=0.12)
                self.shadow(card, blur=20, dist=5, alpha=0.30)
                self.text(894, 42, 224, 22, key_label or "", 15, MUTE)
                self.rich(894, 68, 224, 56,
                          [(key, 52, NAVY, True), (key_unit, 26, NAVY, True)])
        else:
            self.text(M, 44, tw, 52, title, 40, INK, bold=True)
            self.text(M, 96, tw, 28, sub, 19, MUTE)
            self.rule(M, 142, INNER)

    def footer(self, lead, note="メーカー公表値からの計算です（実測ではありません）"):
        self.rule(M, 556, INNER)
        self.text(M, 570, 760, 26, lead, 19, INK, bold=True)
        self.text(M, 596, 800, 22, note, 14, MUTE)
        self.text(836, 596, 300, 22, "寸法で選ぶ　sunpou.nexeed-lab.com", 14, MUTE,
                  align=PP_ALIGN.RIGHT)

    def rows(self, n, top=168, rh=76, highlight=None):
        """等間隔の行。帯を敷いて目で追えるようにする。戻り値は各行の y。"""
        ys = []
        for i in range(n):
            y = top + i * rh
            hot = highlight is not None and i == highlight
            self.rect(52, y, W - 104, rh - 8, HILI if hot else (BAND if i % 2 == 0 else BG))
            if hot:
                self.rect(52, y, 5, rh - 8, NAVY)
            ys.append(y)
        return ys
