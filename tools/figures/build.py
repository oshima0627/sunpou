# -*- coding: utf-8 -*-
"""図版を PowerPoint で作り、PNG に書き出す。

```
python tools/figures/build.py            # 全部
python tools/figures/build.py 名前 名前   # 指定したものだけ
```

出るもの:
  site/public/img/figures.pptx   … 編集できるマスター（1スライド = 1図版）
  site/public/img/<name>.png     … 記事が読む画像（2400x1260 で焼いて等倍表示）

⚠️ **PNG は生成物。直接いじらない。** 直すときは figures.py を直して焼き直す。
⚠️ **pptx を PowerPoint で編集したら、`--from-pptx` で PNG だけ焼き直せる**
   （figures.py からは作り直さない。手で入れた変更が消えるため）。
"""
import argparse
import glob
import json
import os
import re
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from PIL import Image
from pptx import Presentation

import deck
import figures as figmod

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
IMG = os.path.join(ROOT, "site", "public", "img")
PPTX = os.path.join(IMG, "figures.pptx")
SCALE = 2
COLORS = 256

SOFFICE = [
    r"C:/Program Files/LibreOffice/program/soffice.exe",
    r"C:/Program Files (x86)/LibreOffice/program/soffice.exe",
    "soffice",
]

sys.stdout.reconfigure(encoding="utf-8")


def find_soffice():
    for p in SOFFICE:
        if os.path.isfile(p):
            return p
    from shutil import which
    p = which("soffice")
    if p:
        return p
    raise SystemExit("LibreOffice が見つかりません（pptx → PNG の書き出しに要ります）")


def svg_numbers(name):
    """元の SVG に出ている数値トークン。転記ミスの照合に使う。"""
    p = os.path.join(IMG, name + ".svg")
    if not os.path.isfile(p):
        return None
    s = open(p, encoding="utf-8").read()
    body = re.sub(r"<style[\s\S]*?</style>", "", s)
    texts = " ".join(re.findall(r"<text[^>]*>([\s\S]*?)</text>", body))
    texts += " " + " ".join(re.findall(r'aria-label="([^"]*)"', s))
    return set(re.findall(r"\d+(?:\.\d+)?", texts))


def check_numbers(name, used):
    """図に書いた数値が、元の SVG にも出ているか。

    **数字を手で打ち直しているので、転記ミスを機械で捕まえる。**
    元に無い数字が出たら、それは新しく作った数字（＝根拠がない）か打ち間違い。
    """
    src = svg_numbers(name)
    if src is None:
        return []
    return sorted(n for n in used if n not in src)


def check_layout(f):
    """文字同士の重なりと、キャンバスからのはみ出しを総当たりで見る。

    ⚠️ **目視では見落とす。** CLAUDE.md が SVG に対して `getBBox()` を要求しているのと
    同じ理由で、PowerPoint 版でも機械で当たり判定を取る。幅は est_width の近似なので、
    **重なり2px までは許容**する（隣接する行の見かけ上の接触を拾わないため）。
    """
    out = []
    bs = f.boxes
    for i, (x, y, w, h, t) in enumerate(bs):
        if x < 0 or y < 0 or x + w > deck.W + 1 or y + h > deck.H + 1:
            out.append(f"はみ出し「{t[:18]}」 x={x:.0f}..{x + w:.0f} y={y:.0f}..{y + h:.0f}")
        for j in range(i + 1, len(bs)):
            X, Y, Wd, Ht, T = bs[j]
            ox = min(x + w, X + Wd) - max(x, X)
            oy = min(y + h, Y + Ht) - max(y, Y)
            if ox > 2 and oy > 2:
                out.append(f"重なり「{t[:14]}」×「{T[:14]}」 {ox:.0f}x{oy:.0f}px")
    return out


def render(soffice, pptx, outdir):
    filt = ('png:impress_png_Export:{"PixelWidth":{"type":"long","value":%d},'
            '"PixelHeight":{"type":"long","value":%d}}' % (deck.W * SCALE, deck.H * SCALE))
    subprocess.run([soffice, "--headless", "--norestore", "--convert-to", filt,
                    "--outdir", outdir, pptx], capture_output=True, timeout=300)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("names", nargs="*")
    ap.add_argument("--from-pptx", action="store_true",
                    help="figures.py から作り直さず、既存の figures.pptx から PNG だけ焼く")
    a = ap.parse_args()

    soffice = find_soffice()
    order = figmod.ORDER
    if a.names:
        unknown = [n for n in a.names if n not in order]
        if unknown:
            raise SystemExit(f"知らない図です: {unknown}\n使えるのは: {', '.join(order)}")

    if not a.from_pptx:
        prs = Presentation()
        prs.slide_width, prs.slide_height = deck.PX(deck.W), deck.PX(deck.H)
        blank = prs.slide_layouts[6]
        problems = []
        for name in order:
            f = deck.Fig(prs.slides.add_slide(blank))
            figmod.FIGURES[name](f)
            used = set(re.findall(r"\d+(?:\.\d+)?", " ".join(f.drawn)))
            bad = check_numbers(name, used)
            if bad:
                problems.append((name, f"元のSVGに無い数値 {bad}"))
            lay = check_layout(f)
            if lay:
                problems.append((name, "; ".join(lay[:3])))
        if problems:
            for n, b in problems:
                print(f"  ✗ {n}: {b}")
            raise SystemExit("figures.py を直してください")
        prs.save(PPTX)
        print(f"pptx: {os.path.relpath(PPTX, ROOT)}（{len(order)}スライド）")

    # 1スライドずつ焼く（LibreOffice の png 書き出しは先頭スライドだけなので分割する）
    src = Presentation(PPTX)
    targets = a.names or order
    print(f"{'name':34}{'PNG':>9}")
    total = 0
    with tempfile.TemporaryDirectory() as td:
        for i, name in enumerate(order):
            if name not in targets:
                continue
            one = Presentation(PPTX)
            for sl in list(one.slides._sldIdLst)[::-1]:
                idx = list(one.slides._sldIdLst).index(sl)
                if idx != i:
                    one.slides._sldIdLst.remove(sl)
            tmp = os.path.join(td, "one.pptx")
            one.save(tmp)
            render(soffice, tmp, td)
            got = os.path.join(td, "one.png")
            if not os.path.isfile(got):
                raise SystemExit(f"{name}: PNG が出ませんでした")
            out = os.path.join(IMG, name + ".png")
            im = Image.open(got).convert("RGB")
            im.quantize(colors=COLORS, method=Image.MEDIANCUT, dither=Image.NONE).save(
                out, optimize=True)
            os.remove(got)
            total += os.path.getsize(out)
            print(f"{name:34}{os.path.getsize(out) // 1024:>7}KB")
    print(f"\n{len(targets)}枚 / 合計 {total // 1024}KB")


if __name__ == "__main__":
    main()
